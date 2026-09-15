"""Bounded, deterministic observations; no dataset-provided code is executed."""

import json
import sqlite3
from itertools import islice
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

from .config import Settings


def load_tables(path: Path, settings: Settings) -> tuple[list[tuple[str, pd.DataFrame]], bool]:
    """Read at most max_rows+1 rows per table; second result marks table truncation."""
    limit = settings.max_rows + 1
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        frame = pd.read_csv(path, sep="\t" if suffix == ".tsv" else ",", nrows=limit)
        return [(path.name, frame)], False
    if suffix in {".jsonl", ".ndjson"}:
        return [(path.name, pd.read_json(path, lines=True, nrows=limit))], False
    if suffix == ".json":
        # JSON arrays cannot be row-streamed here; file-size admission bounds the input.
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, list) or any(not isinstance(row, dict) for row in data):
            raise ValueError("JSON must be an array of record objects; use JSONL for large inputs")
        return [(path.name, pd.DataFrame(data[:limit]))], False
    if suffix == ".parquet":
        with pq.ParquetFile(path) as source:
            batch = next(source.iter_batches(batch_size=limit), None)
            frame = (
                batch.to_pandas()
                if batch is not None
                else pd.DataFrame(columns=source.schema_arrow.names)
            )
        return [(path.name, frame)], False
    if suffix == ".xlsx":
        with ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > settings.max_file_bytes:
                raise ValueError("Uncompressed workbook exceeds the file-size limit")
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            result = []
            for sheet in book.worksheets[: settings.max_tables]:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                if not header:
                    frame = pd.DataFrame()
                else:
                    names = [f"{index}:{value}" for index, value in enumerate(header)]
                    frame = pd.DataFrame(list(islice(rows, limit)), columns=names)
                result.append((sheet.title, frame))
            return result, len(book.worksheets) > settings.max_tables
        finally:
            book.close()
    if suffix in {".sqlite", ".sqlite3", ".db"}:
        connection = sqlite3.connect(path.as_uri() + "?mode=ro", uri=True)
        try:
            connection.execute("PRAGMA query_only = ON")
            connection.execute("PRAGMA trusted_schema = OFF")
            names = connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table' "
                "AND name NOT LIKE 'sqlite_%' AND sql NOT LIKE 'CREATE VIRTUAL%' "
                "ORDER BY name LIMIT ?",
                (settings.max_tables + 1,),
            ).fetchall()
            result = []
            for (name,) in names[: settings.max_tables]:
                quoted = '"' + name.replace('"', '""') + '"'
                frame = pd.read_sql_query(
                    f"SELECT * FROM {quoted} LIMIT ?", connection, params=(limit,)
                )
                result.append((name, frame))
            return result, len(names) > settings.max_tables
        finally:
            connection.close()
    raise ValueError("Supported formats: CSV, TSV, JSON records, JSONL, Parquet, XLSX, SQLite")


def scalar_key(value):
    if isinstance(value, (dict, list, tuple, np.ndarray)):
        return json.dumps(
            value.tolist() if isinstance(value, np.ndarray) else value,
            sort_keys=True,
            ensure_ascii=False,
            default=str,
        )
    return value


def summarize(name: str, raw: pd.DataFrame, settings: Settings) -> dict:
    frame = raw.iloc[: settings.max_rows, : settings.max_columns]
    columns = []
    normalized = frame.map(scalar_key)
    for index in range(frame.shape[1]):
        series = frame.iloc[:, index]
        stats = {
            "name": str(frame.columns[index])[:200],
            "dtype": str(series.dtype),
            "missing": int(series.isna().sum()),
            "distinct_non_null": int(normalized.iloc[:, index].nunique()),
        }
        if pd.api.types.is_numeric_dtype(series.dtype) and not series.empty:
            for key, value in {
                "min": series.min(),
                "max": series.max(),
                "mean": series.mean(),
            }.items():
                # Convert NaN/Infinity to null for portable JSON.
                number = float(value) if pd.notna(value) else None
                stats[key] = number if number is not None and abs(number) != float("inf") else None
        columns.append(stats)
    return {
        "table": name,
        "rows_analyzed": len(frame),
        "rows_truncated": len(raw) > settings.max_rows,
        "columns_in_source": len(raw.columns),
        "columns_truncated": len(raw.columns) > settings.max_columns,
        "duplicate_rows_in_analyzed_columns": int(normalized.duplicated().sum()),
        "columns": columns,
    }


def profile(path: Path, settings: Settings) -> dict:
    tables, truncated = load_tables(path, settings)
    return {
        "file": path.name,
        "format": path.suffix.lower(),
        "file_bytes": path.stat().st_size,
        "sampling": "First rows only; statistics describe the analyzed subset, not a random sample.",
        "tables_truncated": truncated,
        "tables": [summarize(name, frame, settings) for name, frame in tables],
    }
